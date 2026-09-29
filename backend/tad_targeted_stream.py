
"""Targeted streaming importer for large TAD XLSX workbooks.

This is the transplant from the older Tarrant County research workflow:
scan large county data incrementally, keep only rows that can plausibly match
already-staged leads, and persist those rows in small database batches.
"""

from __future__ import annotations

import json
import re
import uuid
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, Mapping, Optional

from openpyxl import load_workbook

from address_utils import canonical_street_key
from import_staging import stage_import_rows


_ACCOUNT_NAMES = {
    "account", "accountid", "accountnum", "accountnumber", "acctnum",
    "taxaccount", "taxaccountapn", "apn", "pin", "pinnumber",
    "propertyid", "parcelid", "tadaccount", "tadnumber",
}
_ADDRESS_NAMES = {
    "address", "propertyaddress", "situsaddress", "streetaddress",
    "siteaddress", "fulladdress", "matchedaddress",
}
_LEGAL_NAMES = {
    "legaldescription", "propertylegaldescription", "legal",
}


def _norm_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"", "nan", "none", "null", "n/a"}:
        return ""
    return text


def _canonical_account(value: Any) -> str:
    digits = re.sub(r"\D", "", _clean_text(value))
    return digits.lstrip("0") if digits else ""


def _normalize_legal(value: Any) -> str:
    text = _clean_text(value).upper()
    if not text:
        return ""
    text = text.split(",", 1)[0]
    replacements = {
        r"\bBLOCK\b": "BLK",
        r"\bSECTION\b": "SEC",
        r"\bADDITION\b": "ADDN",
        r"\bSUBDIVISION\b": "SUBD",
    }
    for pattern, replacement in replacements.items():
        text = re.sub(pattern, replacement, text)
    text = re.sub(r"[^A-Z0-9\s-]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _house_number(value: Any) -> str:
    match = re.match(r"^\s*(\d+)", _clean_text(value))
    return match.group(1) if match else ""


def _mapping_sources(payload: Mapping[str, Any]) -> Iterable[Mapping[str, Any]]:
    yield payload
    for key in ("raw_import_row", "raw", "source_row", "row"):
        nested = payload.get(key)
        if isinstance(nested, Mapping):
            yield nested


def _first(payload: Mapping[str, Any], names: set[str]) -> Any:
    for source in _mapping_sources(payload):
        for key, value in source.items():
            if _norm_header(key) in names and _clean_text(value):
                return value
    return None


def _decode_payload(value: Any) -> Optional[Dict[str, Any]]:
    if isinstance(value, Mapping):
        return dict(value)
    if isinstance(value, str):
        try:
            decoded = json.loads(value)
        except Exception:
            return None
        return dict(decoded) if isinstance(decoded, Mapping) else None
    return None


async def _load_targets(db, limit: int = 100_000) -> Dict[str, Any]:
    """Read staged non-TAD rows and build compact match indexes."""
    pool = await db.connect()
    try:
        async with pool.acquire() as conn:
            records = await conn.fetch(
                """
                SELECT payload
                FROM import_staging
                WHERE status = 'staged'
                  AND COALESCE(category, '') <> 'tad'
                ORDER BY created_at DESC
                LIMIT $1
                """,
                limit,
            )
    except Exception:
        records = []

    accounts: set[str] = set()
    addresses: set[str] = set()
    legals: set[str] = set()
    by_house: Dict[str, list[str]] = {}

    for record in records:
        payload = _decode_payload(record["payload"])
        if not payload:
            continue

        account = _canonical_account(_first(payload, _ACCOUNT_NAMES))
        if account:
            accounts.add(account)

        address_raw = _clean_text(_first(payload, _ADDRESS_NAMES))
        if address_raw:
            address_key = canonical_street_key(address_raw)
            if address_key:
                addresses.add(address_key)
                house = _house_number(address_raw)
                if house:
                    bucket = by_house.setdefault(house, [])
                    if address_key not in bucket:
                        bucket.append(address_key)

        legal = _normalize_legal(_first(payload, _LEGAL_NAMES))
        if legal:
            legals.add(legal)

    return {
        "accounts": accounts,
        "addresses": addresses,
        "legals": legals,
        "by_house": by_house,
        "count": len(records),
    }


def _header_quality(values: Iterable[Any]) -> int:
    names = {_norm_header(value) for value in values if _clean_text(value)}
    account = bool(names & _ACCOUNT_NAMES)
    location = bool(names & (_ADDRESS_NAMES | _LEGAL_NAMES))
    return int(account) + int(location)


def _row_dict(headers: list[str], values: Iterable[Any]) -> Dict[str, Any]:
    row: Dict[str, Any] = {}
    for index, value in enumerate(values):
        if index >= len(headers):
            break
        header = headers[index]
        if not header:
            continue
        row[header] = value
    return row


def _row_matches(row: Mapping[str, Any], targets: Mapping[str, Any]) -> bool:
    account = _canonical_account(_first(row, _ACCOUNT_NAMES))
    if account and account in targets["accounts"]:
        return True

    legal = _normalize_legal(_first(row, _LEGAL_NAMES))
    if legal and legal in targets["legals"]:
        return True

    address_raw = _clean_text(_first(row, _ADDRESS_NAMES))
    if not address_raw:
        return False

    address_key = canonical_street_key(address_raw)
    if address_key and address_key in targets["addresses"]:
        return True

    # Kidney behavior: use the house number as the cheap first-pass filter, then
    # compare only against lead addresses sharing that house number.
    house = _house_number(address_raw)
    if not house:
        return False

    candidates = targets["by_house"].get(house) or ()
    if not candidates or not address_key:
        return False

    for candidate in candidates:
        if SequenceMatcher(None, address_key, candidate).ratio() >= 0.88:
            return True
    return False


async def stage_targeted_tad_xlsx(
    db,
    workbook_path: Path,
    *,
    source_file: str,
    prepare_row: Callable[[Mapping[str, Any], list[str]], Dict[str, Any]],
    batch_size: int = 1_000,
) -> Dict[str, Any]:
    """Stream a large TAD workbook and stage only rows relevant to known leads."""
    targets = await _load_targets(db)
    if not (targets["accounts"] or targets["addresses"] or targets["legals"]):
        raise ValueError(
            "No staged lead targets are available yet. Upload the foreclosure/distress lead file first."
        )

    batch_id = str(uuid.uuid4())
    categories = ["tad"]
    workbook = load_workbook(
        filename=str(workbook_path),
        read_only=True,
        data_only=True,
    )

    scanned = 0
    matched = 0
    staged = 0
    sheet_reports: list[dict[str, Any]] = []

    try:
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            headers: Optional[list[str]] = None
            inspected = 0

            for values in iterator:
                inspected += 1
                if _header_quality(values) >= 2:
                    headers = [_clean_text(value) for value in values]
                    break
                if inspected >= 50:
                    break

            if not headers:
                sheet_reports.append({
                    "sheet": sheet.title,
                    "rows": 0,
                    "matched": 0,
                    "status": "skipped_no_tad_header",
                })
                continue

            pending: list[Dict[str, Any]] = []
            sheet_scanned = 0
            sheet_matched = 0

            async def flush() -> None:
                nonlocal staged, pending
                if not pending:
                    return
                staged += await stage_import_rows(
                    db,
                    batch_id=batch_id,
                    source_file=source_file,
                    categories=categories,
                    rows=pending,
                )
                pending = []

            for values in iterator:
                if not values or not any(_clean_text(value) for value in values):
                    continue
                row = _row_dict(headers, values)
                scanned += 1
                sheet_scanned += 1

                if not _row_matches(row, targets):
                    continue

                matched += 1
                sheet_matched += 1
                pending.append(prepare_row(row, categories))

                if len(pending) >= batch_size:
                    await flush()

            await flush()
            sheet_reports.append({
                "sheet": sheet.title,
                "rows": sheet_scanned,
                "matched": sheet_matched,
                "status": "ok",
            })
    finally:
        workbook.close()

    return {
        "ok": True,
        "batch_id": batch_id,
        "filename": source_file,
        "categories": categories,
        "files": [{
            "file": source_file,
            "status": "staged",
            "categories": categories,
            "rows": scanned,
            "accepted": staged,
            "rejected": 0,
            "inserted": staged,
            "updated": 0,
            "staged": staged,
            "sheets": sheet_reports,
        }],
        "rows_read": scanned,
        "accepted": staged,
        "rejected": 0,
        "inserted": staged,
        "updated": 0,
        "staged": staged,
        "filtered_out": max(0, scanned - matched),
        "property_ids": [],
        "processing": {
            "status": "staged",
            "matching": "deferred",
            "owner_classification": "deferred",
            "enrichment": "deferred",
        },
        "enrichment": {
            "county": {
                "live_checked": 0,
                "enriched": 0,
                "tad_lookups": 0,
                "missing": 0,
                "deferred": True,
            }
        },
        "transport": {
            "mode": "chunked",
            "parser": "tad_targeted_stream",
            "targets_loaded": targets["count"],
        },
    }
